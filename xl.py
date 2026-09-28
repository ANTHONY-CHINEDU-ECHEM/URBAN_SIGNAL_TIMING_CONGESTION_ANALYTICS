"""Presentation layer shared by all four workbooks."""
import datetime as dt
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule, DataBarRule

NAVY = "1F3864"
INK = "1B2733"
TEAL = "0F766E"
TEAL_L = "CCEDEA"
AMBER = "B45309"
AMBER_L = "FDECC8"
RED = "B91C1C"
RED_L = "FBD5D5"
GREEN = "15803D"
GREEN_L = "D3F0DC"
GREY = "5B6573"
GREY_L = "F3F5F8"
GREY_M = "D8DDE4"
WHITE = "FFFFFF"
INPUT_FILL = "FFF7D6"
FONT = "Arial"

thin = Side(style="thin", color=GREY_M)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
BOTTOM = Border(bottom=Side(style="medium", color=NAVY))


def F(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def setup(ws, title, subtitle, width_last=12, tab=NAVY, cols=14, back=True):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    ws.column_dimensions["A"].width = 2.5
    for c in range(2, cols + 2):
        ws.column_dimensions[L(c)].width = width_last
    for c in range(1, cols + 3):
        ws.cell(1, c).fill = fill(NAVY)
        ws.cell(2, c).fill = fill(NAVY)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 20
    ws["B1"] = title
    ws["B1"].font = F(17, True, WHITE)
    ws["B1"].alignment = Alignment(vertical="center")
    ws["B2"] = subtitle
    ws["B2"].font = F(10, False, "C9D3E3")
    ws["B2"].alignment = Alignment(vertical="top")
    if back:
        c = ws.cell(3, cols + 1)
        c.value = "Back to Cover"
        c.hyperlink = "#'Cover'!A1"
        c.font = Font(name=FONT, size=9, color=TEAL, underline="single")
        c.alignment = Alignment(horizontal="right")
    ws.freeze_panes = "A4"


def section(ws, row, col, text, span=8):
    c = ws.cell(row, col, text)
    c.font = F(12, True, NAVY)
    for k in range(col, col + span):
        ws.cell(row, k).border = BOTTOM
    ws.row_dimensions[row].height = 20
    return row + 1


def note(ws, row, col, text, span=10, height=None, size=9, color=GREY, italic=False):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
    c = ws.cell(row, col, text)
    c.font = F(size, False, color, italic)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if height is None:
        est_chars = sum(ws.column_dimensions[L(k)].width or 9 for k in range(col, col + span)) * 0.95 * (9 / size)
        lines = max(1, int(len(text) / max(20, est_chars)) + 1 + text.count("\n"))
        height = 13.5 * lines + 3
    ws.row_dimensions[row].height = height
    return row + 1


def kv(ws, row, col, label, value, fmt=None, is_input=False, span_label=2, comment=None, bold=False):
    ws.cell(row, col, label).font = F(10, bold, INK)
    for k in range(span_label - 1):
        pass
    c = ws.cell(row, col + span_label, value)
    c.font = F(10, True if bold else False, "0000FF" if is_input else INK)
    c.border = BORDER
    if is_input:
        c.fill = fill(INPUT_FILL)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal="right")
    if comment:
        from openpyxl.comments import Comment
        c.comment = Comment(comment, "Model")
    return c


def header_row(ws, row, col, headers, fillc=NAVY, height=30):
    for i, h in enumerate(headers):
        c = ws.cell(row, col + i, h)
        c.font = F(9, True, WHITE)
        c.fill = fill(fillc)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = height


def body_cell(c, fmt=None, zebra=False, bold=False, color=INK, align=None):
    c.font = F(9, bold, color)
    c.border = BORDER
    if zebra:
        c.fill = fill(GREY_L)
    if fmt:
        c.number_format = fmt
    if align and align != "wrap":
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    elif align == "wrap":
        c.alignment = Alignment(wrap_text=True, vertical="center")
    else:
        c.alignment = Alignment(vertical="center")


def write_table(ws, row, col, df, fmts=None, wrap_cols=(), max_rows=None, zebra=True, header_height=30):
    """Write a DataFrame as a styled table. Values beginning with '=' are formulas."""
    fmts = fmts or {}
    header_row(ws, row, col, list(df.columns), height=header_height)
    data = df if max_rows is None else df.head(max_rows)
    for i, rec in enumerate(data.itertuples(index=False), start=1):
        for j, v in enumerate(rec):
            if isinstance(v, (np.floating,)):
                v = None if np.isnan(v) else float(v)
            elif isinstance(v, np.integer):
                v = int(v)
            elif isinstance(v, np.bool_):
                v = bool(v)
            elif isinstance(v, pd.Timestamp):
                v = None if pd.isna(v) else v.to_pydatetime()
            elif v is pd.NaT:
                v = None
            c = ws.cell(row + i, col + j, v)
            name = df.columns[j]
            body_cell(c, fmts.get(name), zebra and i % 2 == 0,
                      align="wrap" if name in wrap_cols else None)
    return row + len(data) + 1


def kpi(ws, row, col, label, formula, fmt, sub="", width=3, accent=TEAL):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + width - 1)
    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + width - 1)
    ws.merge_cells(start_row=row + 2, start_column=col, end_row=row + 2, end_column=col + width - 1)
    for r in range(row, row + 3):
        for k in range(col, col + width):
            ws.cell(r, k).fill = fill(WHITE)
            ws.cell(r, k).border = Border(
                left=Side(style="thick", color=accent) if k == col else None,
                top=thin if r == row else None, bottom=thin if r == row + 2 else None,
                right=thin if k == col + width - 1 else None)
    a = ws.cell(row, col, label)
    a.font = F(8, True, GREY)
    a.alignment = Alignment(horizontal="left", indent=1, vertical="bottom")
    b = ws.cell(row + 1, col, formula)
    b.font = F(18, True, NAVY)
    b.number_format = fmt
    b.alignment = Alignment(horizontal="left", indent=1, vertical="center")
    s = ws.cell(row + 2, col, sub)
    s.font = F(8, False, GREY, True)
    s.alignment = Alignment(horizontal="left", indent=1, vertical="top")
    ws.row_dimensions[row + 1].height = 30


def selector(ws, cell, source, default, label=None, label_cell=None, prompt=None):
    dv = DataValidation(type="list", formula1=source, allow_blank=False, showDropDown=False)
    dv.promptTitle = "Selector"
    dv.prompt = prompt or "Choose a value from the drop down list"
    dv.showInputMessage = True
    dv.error = "Please choose a value from the list"
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(cell)
    c = ws[cell]
    c.value = default
    c.font = F(11, True, NAVY)
    c.fill = fill(WHITE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    med = Side(style="medium", color=TEAL)
    c.border = Border(left=med, right=med, top=med, bottom=med)
    if label and label_cell:
        lc = ws[label_cell]
        lc.value = label
        lc.font = F(8, True, GREY)
        lc.alignment = Alignment(horizontal="left", vertical="bottom")


def add_name(wb, name, ref):
    dn = DefinedName(name, attr_text=ref)
    wb.defined_names[name] = dn


def input_style(c, fmt=None):
    c.font = F(10, False, "0000FF")
    c.fill = fill(INPUT_FILL)
    c.border = BORDER
    if fmt:
        c.number_format = fmt


def calc_style(c, fmt=None, bold=False, color=INK):
    c.font = F(10, bold, color)
    c.border = BORDER
    if fmt:
        c.number_format = fmt


def link_style(c, fmt=None):
    c.font = F(10, False, "008000")
    c.border = BORDER
    if fmt:
        c.number_format = fmt


def dump_df(ws, df, start_row=1, fmts=None, header_fill=NAVY):
    """Fast bulk dump of a large DataFrame (data sheets)."""
    fmts = fmts or {}
    cols = list(df.columns)
    for j, h in enumerate(cols, 1):
        c = ws.cell(start_row, j, h)
        c.font = F(9, True, WHITE)
        c.fill = fill(header_fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[start_row].height = 32
    conv = []
    for c in cols:
        s = df[c]
        if pd.api.types.is_datetime64_any_dtype(s):
            conv.append([None if pd.isna(x) else x.to_pydatetime() for x in s])
        elif s.dtype == bool:
            conv.append([bool(x) for x in s])
        else:
            conv.append([None if (x is None or (isinstance(x, float) and np.isnan(x)) or x is pd.NaT) else
                         (int(x) if isinstance(x, np.integer) else float(x) if isinstance(x, np.floating) else
                          bool(x) if isinstance(x, np.bool_) else x) for x in s])
    for i in range(len(df)):
        ws.append([conv[j][i] for j in range(len(cols))]) if start_row == 1 and i >= 0 else None
    for j, c in enumerate(cols, 1):
        if c in fmts:
            for r in range(start_row + 1, start_row + len(df) + 1):
                ws.cell(r, j).number_format = fmts[c]
        ws.column_dimensions[L(j)].width = max(10, min(30, len(c) + 2))
    ws.freeze_panes = ws.cell(start_row + 1, 2)
    ws.auto_filter.ref = f"A{start_row}:{L(len(cols))}{start_row + len(df)}"


def fit_text_width(ws, col, width):
    ws.column_dimensions[L(col)].width = width


def pct_fmt():
    return '0.0%;(0.0%);"-"'


def block_note(ws, r1, c1, r2, c2, text, size=9, color=GREY, bg=GREY_L, title=None):
    """Merged multi row explanatory panel that does not stretch row heights."""
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    c = ws.cell(r1, c1, (title + "\n" if title else "") + text)
    c.font = F(size, False, color)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    for r in range(r1, r2 + 1):
        for k in range(c1, c2 + 1):
            ws.cell(r, k).fill = fill(bg)


def print_fit(ws, area, landscape=True, tall=1):
    ws.print_area = area
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = tall
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def row_h(text, chars_per_line, base=13.5):
    lines = max(1, -(-len(str(text)) // chars_per_line))
    return base * lines + 4


def ctitle(chart, text, size=1100, color=NAVY):
    from openpyxl.chart.title import Title
    from openpyxl.chart.text import RichText, Text
    from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, RegularTextRun
    cp = CharacterProperties(sz=size, b=True, solidFill=color, latin=None)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    chart.title = Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def unsmooth(wb):
    """Line series default to smoothed rendering in some viewers; draw straight segments unless smoothing was set deliberately."""
    for ws in wb.worksheets:
        for ch in ws._charts:
            for sub in [ch] + list(getattr(ch, "_charts", [])):
                for s in getattr(sub, "series", []):
                    if s.smooth is None:
                        s.smooth = False
